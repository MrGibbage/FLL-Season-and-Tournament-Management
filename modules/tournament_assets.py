"""Tournament configuration stored with OJS files and shared TOAST assets."""
import copy
import json
from pathlib import Path
import sys

CONFIG_SHEET = "_TOAST_Config"
CONFIG_VERSION = 1
CHUNK_SIZE = 30000  # Below Excel's per-cell text limit.


def write_embedded_config(book, config, division=""):
    """Write identical config payloads, with separate workbook identity metadata."""
    if CONFIG_SHEET in book.sheetnames:
        del book[CONFIG_SHEET]
    sheet = book.create_sheet(CONFIG_SHEET)
    sheet.sheet_state = "veryHidden"
    sheet["A1"], sheet["B1"] = "TOAST_CONFIG", CONFIG_VERSION
    sheet["A2"], sheet["B2"] = "division", division
    info = config["INFO"]
    sheet["A3"], sheet["B3"] = "tournament", json.dumps(
        [info.get("season_year"), info["tournament_short_name"]])
    payload = json.dumps(config, ensure_ascii=True, sort_keys=True)
    chunks = [payload[i:i + CHUNK_SIZE] for i in range(0, len(payload), CHUNK_SIZE)]
    sheet["A4"], sheet["B4"] = "chunks", len(chunks)
    for row, chunk in enumerate(chunks, 5):
        sheet.cell(row, 1, chunk)


def read_embedded_config(book):
    if CONFIG_SHEET not in book.sheetnames:
        return None
    sheet = book[CONFIG_SHEET]
    if sheet["A1"].value != "TOAST_CONFIG" or sheet["B1"].value != CONFIG_VERSION:
        raise ValueError("Unsupported embedded TOAST configuration version")
    division = sheet["B2"].value or ""
    if division not in ("", "D1", "D2"):
        raise ValueError(f"Invalid embedded division: {division}")
    count = sheet["B4"].value
    if type(count) is not int or not 1 <= count <= 1000:
        raise ValueError("Invalid embedded configuration chunk count")
    chunks = [sheet.cell(row, 1).value for row in range(5, 5 + count)]
    if any(not isinstance(chunk, str) for chunk in chunks):
        raise ValueError("Incomplete embedded TOAST configuration")
    config = json.loads("".join(chunks))
    if not isinstance(config, dict) or not isinstance(config.get("INFO"), dict) or not isinstance(config.get("AWARDS"), list):
        raise ValueError("Invalid embedded TOAST configuration")
    return division, sheet["B3"].value, config


def discover_embedded_config(directory, workbooks):
    """Use the lowest division present, even when other payloads disagree."""
    records = []
    for path in sorted(Path(directory).iterdir()):
        if not path.is_file() or path.suffix.lower() != ".xlsm" or path.name.startswith("~$"):
            continue
        record = read_embedded_config(workbooks.workbook(str(path)))
        if record is not None:
            division, identity, config = record
            records.append((division, identity, config, path))
    if not records:
        return None
    records.sort(key=lambda record: record[0])
    divisions = [r[0] for r in records]
    if len(set(divisions)) != len(divisions) or ("" in divisions and len(records) > 1):
        raise ValueError("Multiple OJS files for the same division; use one tournament per folder")
    if len({r[1] for r in records}) != 1:
        raise ValueError("OJS files from different tournaments; use one tournament per folder")
    division, _, authoritative, path = records[0]
    config = copy.deepcopy(authoritative)
    info = config["INFO"]
    if info.get("using_divisions") != bool(division):
        raise ValueError("Embedded division identity disagrees with authoritative configuration")
    expected = {entry["division"] for entry in info["ojs_files"]}
    if expected != set(divisions):
        raise ValueError(f"Missing or unexpected OJS divisions: expected {sorted(expected)}, found {divisions}")
    warnings = []
    for other_division, _, other, other_path in records[1:]:
        if other != authoritative:
            warnings.append(f"{other_path.name}: configuration differs; using {path.name} ({division})")
    # Discover by workbook identity, so renaming an OJS does not break TOAST.
    info["ojs_files"] = [{"division": r[0], "filename": r[3].name} for r in records]
    return config, warnings, path.name


def tournament_display_title(info):
    """Label the whole event without repeating a single division on each award."""
    if info.get("using_divisions"):
        divisions = sorted({entry.get("division", "") for entry in info.get("ojs_files", [])})
        numbers = [division[1:] for division in divisions if division in ("D1", "D2")]
        label = ("Divisions " if len(numbers) > 1 else "Division ") + " & ".join(numbers)
        if not numbers:
            label = "Divisions 1 & 2"
    else:
        label = info.get("division_label") or "Single Division"
    return f"{info.get('tournament_long_name', '')} - {label}"


def toast_template_directory():
    """Source runs read editable project files; frozen runs read bundled files."""
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parents[1]


def toast_template_name(kind, using_divisions):
    suffix = "with-divisions" if using_divisions else "without-divisions"
    return f"{kind}_template-{suffix}.html.jinja"


def finalize_tournament_assets(folder):
    """Embed the complete merged config only after Maestro finishes all divisions."""
    from openpyxl import load_workbook
    folder = Path(folder)
    config_path = folder / "tournament_config.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    for entry in config["INFO"]["ojs_files"]:
        path = folder / entry["filename"]
        if not path.is_file() or path.with_name("~$" + path.name).exists():
            raise ValueError(f"Missing or open OJS workbook: {path}")
    for entry in config["INFO"]["ojs_files"]:
        path = folder / entry["filename"]
        book = load_workbook(path, keep_vba=True)
        try:
            write_embedded_config(book, config, entry.get("division", ""))
            book.save(path)
        finally:
            book.close()
            if book.vba_archive is not None:
                book.vba_archive.close()
    # Keep a local authoring copy outside the director's tournament folder.
    archive = folder.parent / "_maestro-config"
    archive.mkdir(exist_ok=True)
    (archive / f"{folder.name}.json").write_text(config_path.read_text(encoding="utf-8"), encoding="utf-8")
    config_path.unlink()
