"""Read-only, per-run cache of Excel's saved values and parsed tables."""

import os
import pandas as pd
from openpyxl import load_workbook


class OJSWorkbookUnavailable(ValueError):
    """An expected workbook access problem with actionable director instructions."""
    def __init__(self, path):
        self.path = path
        super().__init__(
            f"TOAST cannot read the OJS workbook: {os.path.basename(path)}\n"
            "It may be open in Excel or locked by another application.\n"
            "Save and close the workbook in Excel, then run TOAST again.\n"
            "If it is already closed, check that you have permission to read the file."
        )


class CeremonyWorkbooks:
    def __init__(self):
        self.books = {}
        self.tables = {}

    def workbook(self, path):
        path = os.path.abspath(path)
        if path not in self.books:
            lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
            if os.path.exists(lock):
                raise OJSWorkbookUnavailable(path)
            try:
                self.books[path] = load_workbook(path, data_only=True)
            except PermissionError as error:
                raise OJSWorkbookUnavailable(path) from error
        return self.books[path]

    def table(self, path, sheet, name):
        key = (os.path.abspath(path), sheet, name)
        if key not in self.tables:
            ws = self.workbook(path)[sheet]
            rows = list(ws[ws.tables[name].ref])
            headers = [str(c.value).strip() for c in rows[0]]
            if len(headers) != len(set(headers)):
                raise ValueError(f"Duplicate headers in {name}")
            values = [[c.value.strip() if isinstance(c.value, str) else c.value
                       for c in row] for row in rows[1:]]
            self.tables[key] = pd.DataFrame(values, columns=headers).replace('', None)
        return self.tables[key].copy()

    def close(self):
        for book in self.books.values():
            book.close()
        self.books.clear()
        self.tables.clear()
