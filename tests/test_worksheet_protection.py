"""Protection follows resized OJS table bounds, not template formatting."""
import unittest
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Protection

from modules.worksheet_setup import protect_worksheets


ROOT = Path(__file__).resolve().parents[1]


class WorksheetProtectionTests(unittest.TestCase):
    def test_only_table_data_cells_can_remain_unlocked(self):
        for template in ROOT.glob('2026-Qualifier-Template-*.xlsm'):
            with self.subTest(template=template.name):
                book = openpyxl.load_workbook(template, keep_vba=True)
                try:
                    sheet = book['Team and Program Information']
                    sheet.tables['OfficialTeamList'].ref = 'A2:D6'
                    sheet['D4'].protection = Protection(locked=False)
                    sheet['A7'].protection = Protection(locked=False)
                    sheet['E4'].protection = Protection(locked=False)
                    protect_worksheets(pd.Series({'OJS_FileName': template.name}), book)
                    self.assertFalse(sheet['D4'].protection.locked)
                    self.assertFalse(sheet['F2'].protection.locked)
                    self.assertTrue(sheet['A7'].protection.locked)
                    self.assertTrue(sheet['E4'].protection.locked)
                    self.assertTrue(sheet.protection.sheet)
                    results = book['Results and Rankings']
                    self.assertTrue(results.column_dimensions['A'].protection.locked)
                    self.assertTrue(results.column_dimensions['Z'].protection.locked)
                    self.assertFalse(results['P3'].protection.locked)
                    with tempfile.TemporaryDirectory() as temp_dir:
                        output = Path(temp_dir) / 'protected.xlsm'
                        book.save(output)
                        saved = openpyxl.load_workbook(output, keep_vba=True)
                        try:
                            saved_results = saved['Results and Rankings']
                            self.assertTrue(saved_results.column_dimensions['A'].protection.locked)
                            self.assertTrue(saved_results.column_dimensions['Z'].protection.locked)
                            self.assertFalse(saved_results['P3'].protection.locked)
                            self.assertFalse(saved['Team and Program Information']['F2'].protection.locked)
                        finally:
                            saved.close()
                    for ws in book.worksheets:
                        self.assertTrue(all(dim.protection.locked
                                            for dim in ws.column_dimensions.values()))
                        self.assertTrue(all(dim.protection.locked
                                            for dim in ws.row_dimensions.values()))
                        areas = [openpyxl.worksheet.cell_range.CellRange(t.ref)
                                 for t in ws.tables.values()]
                        for cell in ws._cells.values():
                            if not cell.protection.locked:
                                if ws.title == 'Team and Program Information' and cell.coordinate == 'F2':
                                    continue
                                self.assertTrue(any(a.min_col <= cell.column <= a.max_col
                                                    and a.min_row < cell.row <= a.max_row
                                                    for a in areas), f'{ws.title}!{cell.coordinate}')
                finally:
                    book.close()


if __name__ == '__main__':
    unittest.main()
