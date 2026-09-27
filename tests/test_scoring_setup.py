"""Regression checks using the current season fixtures; no workbooks are written."""

import io
from pathlib import Path
import unittest
import warnings

import openpyxl

from modules.excel_operations import read_table_as_df
from modules.worksheet_setup import (
    copy_award_def,
    resize_worksheets,
    set_up_tapi_worksheet,
)


ROOT = Path(__file__).resolve().parents[1]


class ScoringSetupTests(unittest.TestCase):
    def test_generated_scoring_and_award_definitions(self):
        for divisions, name, expected_counts in [
            (True, "Dunford", [1, 2]),
            (False, "Cinderbrook", [3]),
        ]:
            suffix = "with-divisions" if divisions else "without-divisions"
            source = ROOT / f"2025-FLL-Qualifier-Tournaments-{suffix}.xlsx"
            sheet, table = ("DivTournaments", "DivTournamentList") if divisions else (
                "Tournaments", "TournamentList"
            )
            tournaments = read_table_as_df(str(source), sheet, table).fillna(0)
            awards = read_table_as_df(str(source), "AwardDef", "AwardDef").fillna(0)
            assignments = read_table_as_df(str(source), "Assignments", "Assignments").fillna(0)
            counts = []
            for _, tournament in tournaments[tournaments["Short Name"] == name].iterrows():
                with self.subTest(tournament=name, division=tournament.get("Div")):
                    with warnings.catch_warnings():
                        warnings.simplefilter("ignore", UserWarning)
                        book = openpyxl.load_workbook(
                            ROOT / f"2025-Qualifier-Template-{suffix}.xlsm", keep_vba=True
                        )
                    self.assertTrue(set_up_tapi_worksheet(
                        tournament, book, assignments, divisions
                    ))
                    copy_award_def(tournament, book, awards)
                    resize_worksheets(tournament, book, assignments, divisions)

                    # Round-trip in memory to check the formulas and table metadata
                    # that Excel will receive, while retaining the VBA archive.
                    output = io.BytesIO()
                    book.save(output)
                    book.close()
                    output.seek(0)
                    with warnings.catch_warnings():
                        warnings.simplefilter("ignore", UserWarning)
                        saved = openpyxl.load_workbook(output, keep_vba=True)
                    cv = saved["Core Values Input"]
                    formula = "CoreValuesResults[[#This Row],[Gracious Professionalism Total]]"
                    end = openpyxl.utils.range_boundaries(cv.tables["CoreValuesResults"].ref)[3]
                    self.assertGreater(end, 2)
                    for row in range(2, end + 1):
                        self.assertEqual(cv.cell(row, 18).value, "=" + formula)
                        self.assertIn("[Gracious Professionalism Score]", cv.cell(row, 19).value)
                    column = next(c for c in cv.tables["CoreValuesResults"].tableColumns
                                  if c.name == "Gracious Professionalism Score")
                    self.assertEqual(column.calculatedColumnFormula.attr_text, formula)
                    # Different teams must contribute their own GP total once.
                    for row, scores in [(2, (4, 2, 3)), (3, (2, 2, 3))]:
                        for col, score in zip(range(14, 17), scores):
                            cv.cell(row, col).value = score
                        self.assertIn("[Gracious Professionalism 1]:[Gracious Professionalism 3]",
                                      cv.cell(row, 17).value)
                    self.assertEqual(sum(cv.cell(2, c).value for c in range(14, 17)), 9)
                    self.assertEqual(sum(cv.cell(3, c).value for c in range(14, 17)), 7)

                    ws = saved["AwardDef"]
                    data = list(ws[ws.tables["AwardDef"].ref])
                    entries = {r[0].value: r[3].value for r in data[1:]}
                    self.assertEqual(entries["ADV"], int(tournament["ADV"]))
                    for award_id, count in entries.items():
                        self.assertEqual(count, int(tournament.get(award_id, 0)))
                    counts.append(entries["ADV"])
                    self.assertEqual(saved["Results and Rankings"]["Q1"].value,
                                     '=VLOOKUP("ADV", AwardDef[], 4,0)')
                    self.assertIn("xl/vbaProject.bin", saved.vba_archive.namelist())
                    saved.close()
            self.assertEqual(sorted(counts), expected_counts)


if __name__ == "__main__":
    unittest.main()
