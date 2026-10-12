"""Embedded config, lowest-division authority, and source/bundle template loading."""
import copy
import io
import json
import logging
from pathlib import Path
import runpy
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from modules.ceremony_workbooks import CeremonyWorkbooks
from modules.tournament_assets import (
    CONFIG_SHEET, write_embedded_config, read_embedded_config,
    discover_embedded_config, finalize_tournament_assets, toast_template_directory, tournament_display_title,
)

ROOT = Path(__file__).resolve().parents[1]


def config_for(divisions=("D1", "D2")):
    return {"INFO": {"season_year": "2026", "tournament_short_name": "Test",
                     "using_divisions": divisions != ("",),
                     "ojs_files": [{"division": d, "filename": f"ojs{d}.xlsm"} for d in divisions]},
            "AWARDS": []}


class EmbeddedConfigTests(unittest.TestCase):
    def test_division_titles(self):
        info = config_for()["INFO"]
        info["tournament_long_name"] = "Test Tournament"
        self.assertEqual(tournament_display_title(info), "Test Tournament - Divisions 1 & 2")
        info = config_for(("D2",))["INFO"]
        info["tournament_long_name"] = "Test Tournament"
        self.assertEqual(tournament_display_title(info), "Test Tournament - Division 2")
        info = config_for(("",))["INFO"]
        info["tournament_long_name"] = "Test Tournament"
        self.assertEqual(tournament_display_title(info), "Test Tournament - Single Division")
        info["division_label"] = "Division 2"
        self.assertEqual(tournament_display_title(info), "Test Tournament - Division 2")

    def test_saved_roundtrip_large_unicode_payload(self):
        config = config_for()
        config["INFO"]["message"] = "Robotics & robots < > \u00e9" * 4000
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.xlsx"
            book = Workbook()
            write_embedded_config(book, config, "D1")
            book.save(path)
            book.close()
            book = load_workbook(path, data_only=True)
            try:
                self.assertEqual(read_embedded_config(book)[2], config)
                self.assertEqual(book[CONFIG_SHEET].sheet_state, "veryHidden")
            finally:
                book.close()

    def make_cache(self, folder, config, divisions, changed=None):
        cache = CeremonyWorkbooks()
        for division in divisions:
            path = folder / f"renamed-{division or 'only'}.xlsm"
            path.touch()
            book = Workbook()
            write_embedded_config(book, changed if changed and division == "D2" else config, division)
            cache.books[str(path.resolve())] = book
        self.addCleanup(cache.close)
        return cache

    def test_lowest_division_wins_and_renamed_files_work(self):
        config = config_for()
        changed = copy.deepcopy(config)
        changed["AWARDS"] = [{"ID": "different"}]
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            cache = self.make_cache(folder, config, ["D2", "D1"], changed)
            actual, warnings, source = discover_embedded_config(folder, cache)
            self.assertEqual(actual["AWARDS"], [])
            self.assertEqual(source, "renamed-D1.xlsm")
            self.assertEqual(len(warnings), 1)
            self.assertEqual(actual["INFO"]["ojs_files"][0]["filename"], "renamed-D1.xlsm")

    def test_sole_d2_and_nondivision(self):
        for divisions in [("D2",), ("",)]:
            with self.subTest(divisions=divisions), tempfile.TemporaryDirectory() as directory:
                folder = Path(directory)
                cache = self.make_cache(folder, config_for(divisions), divisions)
                config, warnings, source = discover_embedded_config(folder, cache)
                self.assertEqual(config["INFO"]["ojs_files"][0]["division"], divisions[0])
                self.assertFalse(warnings)

    def test_missing_duplicate_and_mixed_tournaments_rejected(self):
        for problem in ["missing", "duplicate", "mixed", "version"]:
            with self.subTest(problem=problem), tempfile.TemporaryDirectory() as directory:
                folder = Path(directory)
                cache = self.make_cache(folder, config_for(), ["D1", "D2"] if problem != "missing" else ["D1"])
                if problem != "missing":
                    book = cache.workbook(str(folder / "renamed-D2.xlsm"))
                    book[CONFIG_SHEET][{"duplicate": "B2", "mixed": "B3", "version": "B1"}[problem]] = {
                        "duplicate": "D1", "mixed": "other tournament", "version": 999}[problem]
                with self.assertRaises(ValueError):
                    discover_embedded_config(folder, cache)

    def test_final_pass_embeds_complete_config_and_archives_json(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / "Test"
            folder.mkdir()
            config = config_for()
            for entry in config["INFO"]["ojs_files"]:
                book = Workbook()
                book.save(folder / entry["filename"])
                book.close()
            (folder / "tournament_config.json").write_text(json.dumps(config), encoding="utf-8")
            finalize_tournament_assets(folder)
            self.assertFalse((folder / "tournament_config.json").exists())
            self.assertTrue((folder.parent / "_maestro-config/Test.json").exists())
            for entry in config["INFO"]["ojs_files"]:
                book = load_workbook(folder / entry["filename"])
                try:
                    self.assertEqual(read_embedded_config(book)[2], config)
                finally:
                    book.close()

    def test_source_and_frozen_asset_paths(self):
        self.assertEqual(toast_template_directory(), ROOT)
        with patch.object(sys, "frozen", True, create=True), patch.object(sys, "_MEIPASS", "bundled-assets", create=True):
            self.assertEqual(toast_template_directory(), Path("bundled-assets"))

    def test_scored_tournaments_without_loose_assets(self):
        for name in ["Dunford", "Brindlewick"]:
            with self.subTest(tournament=name), tempfile.TemporaryDirectory() as directory:
                source = ROOT / "tournaments" / name
                if not (source / "tournament_config.json").exists():
                    self.skipTest("Requires local scored tournament fixtures")
                config = json.loads((source / "tournament_config.json").read_text(encoding="utf-8"))
                folder = Path(directory)
                cache = CeremonyWorkbooks()
                try:
                    for entry in config["INFO"]["ojs_files"]:
                        path = folder / entry["filename"]
                        shutil.copyfile(source / entry["filename"], path)
                        # Keep the real saved calculations intact; add metadata in memory.
                        write_embedded_config(cache.workbook(str(path)), config, entry["division"])
                    argv = ["toast", "--tournament-dir", str(folder), "--output-dir", str(folder / "out")]
                    # Exercise source loading, then simulate the packaged asset directory.
                    for frozen in (False, True):
                        bundle = folder / "bundle"
                        bundle.mkdir(exist_ok=True)
                        for template in ROOT.glob("*_template-*.html.jinja"):
                            shutil.copyfile(template, bundle / template.name)
                        with patch.object(sys, "frozen", frozen, create=True), \
                             patch.object(sys, "_MEIPASS", str(bundle), create=True), \
                             patch("sys.argv", argv), patch("builtins.input", return_value="y"), \
                             patch("sys.stdout", new=io.StringIO()):
                            runpy.run_path(str(ROOT / "fll-toast.py"))["main"](cache)
                    summary = next((folder / "out").glob("*-summary.html")).read_text(encoding="utf-8")
                    for entry in config["INFO"]["ojs_files"]:
                        df = cache.table(str(folder / entry["filename"]), "Results and Rankings", "TournamentData")
                        for _, row in df[df["Advance?"].isin(["Yes", "Alt"])].iterrows():
                            self.assertIn(f"Team {int(row['Team #'])},", summary)
                    self.assertFalse(list(folder.glob("*.jinja")))
                    self.assertFalse(list(folder.glob("*.json")))
                finally:
                    cache.close()
                    logger = logging.getLogger("ceremony_generator")
                    for handler in list(logger.handlers):
                        handler.close()
                        logger.removeHandler(handler)
