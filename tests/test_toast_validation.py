import copy
import io
import json
import logging
from pathlib import Path
import runpy
import tempfile
import unittest
from unittest.mock import patch
import warnings

import pandas as pd
from openpyxl import load_workbook

from modules.ceremony_workbooks import CeremonyWorkbooks
from modules.ceremony_validator import OJSValidator
from modules.ceremony_data_collector import CeremonyDataCollector
from modules.ceremony_renderer import CeremonyRenderer

ROOT = Path(__file__).resolve().parents[1]


class ScoreBoundaryTests(unittest.TestCase):
    def test_robot_score_boundaries_and_override(self):
        for value, maximum, valid in [(530, 530, True), (531, 530, False),
                                      (531, 600, True), (-1, 530, False),
                                      (2.5, 530, False), (None, 530, False),
                                      ('bad', 530, False), (True, 530, False)]:
            with self.subTest(value=value, maximum=maximum):
                cache = CeremonyWorkbooks()
                with patch.object(cache, 'table', return_value=pd.DataFrame({
                    f'Robot Game {i} Score': [value] for i in range(1, 4)
                })):
                    validator = OJSValidator(cache, maximum)
                    self.assertEqual(validator.validate_robot_game_scores('unused'), valid)

    def test_rubrics_and_gp(self):
        for value in [5, 2.5, -1, None]:
            cache = CeremonyWorkbooks()
            with patch.object(cache, 'table', return_value=pd.DataFrame({'Rubric': [value]})):
                self.assertFalse(OJSValidator(cache).validate_rubric_scores('x', 's', 't', ['Rubric']))
        for value in [0, 2, 3, 4]:
            cache = CeremonyWorkbooks()
            with patch.object(cache, 'table', return_value=pd.DataFrame({
                f'Gracious Professionalism {i}': [value] for i in range(1,4)
            })):
                self.assertTrue(OJSValidator(cache).validate_core_values_scores('x'))

    def test_missing_template_variable_fails(self):
        with tempfile.TemporaryDirectory() as folder:
            Path(folder, 'test.jinja').write_text('{% if missing %}{{ missing }}{% endif %}')
            renderer = CeremonyRenderer(folder)
            self.assertIn('missing', renderer.extract_template_variables('test.jinja'))
            with self.assertRaises(Exception):
                renderer.render_text('test.jinja', {})


class TournamentValidationTests(unittest.TestCase):
    def setUp(self):
        source = ROOT / 'tournaments/Dunford/tournament_config.json'
        if not source.exists():
            self.skipTest('Requires local Dunford fixtures')
        self.config = json.loads(source.read_text())
        self.cache = CeremonyWorkbooks()
        self.addCleanup(self.cache.close)
        self.paths = [str(source.parent / e['filename']) for e in self.config['INFO']['ojs_files']]
        # Correct the known extra Judges 1 selection only in this in-memory fixture.
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            df = self.cache.table(self.paths[1], 'Results and Rankings', 'TournamentData')
        df.loc[df['Award'] == 'Judges 1', 'Award'] = None
        self.cache.tables[(self.paths[1], 'Results and Rankings', 'TournamentData')] = df

    def validate(self):
        validator = OJSValidator(self.cache)
        for path in self.paths:
            validator.validate_all_sheets(path)
        validator.validate_tournament(self.config, str(ROOT / 'tournaments/Dunford'))
        return validator

    def test_valid_scores_and_one_load_per_workbook(self):
        with patch('modules.ceremony_workbooks.load_workbook', wraps=load_workbook) as loader:
            validator = self.validate()
            self.assertFalse(validator.errors, [str(e) for e in validator.errors])
            collector = CeremonyDataCollector(self.config, workbooks=self.cache)
            for path in self.paths:
                collector.collect_team_list(path)
                collector.collect_advancing_teams(path)
                collector.collect_robot_game_awards(path, 1)
            # D2 was already opened in setUp; D1 opens once despite all consumers.
            self.assertEqual(loader.call_count, 1)

    def test_stale_scores_duplicate_awards_and_advancement(self):
        df = self.cache.tables[(self.paths[1], 'Results and Rankings', 'TournamentData')]
        df.loc[0, "Champion's Score"] = 999
        df.loc[0, 'Advance?'] = 'Maybe'
        df.loc[0, 'Award'] = 'Judges 1'
        errors = '\n'.join(str(e) for e in self.validate().errors)
        self.assertIn('saved calculation', errors)
        self.assertIn('Advance?', errors)
        self.assertIn('Judges 1', errors)

    def test_division_render_has_both_team_lists(self):
        module = runpy.run_path(str(ROOT / 'fll-toast.py'))
        with tempfile.TemporaryDirectory() as folder:
            argv = ['fll-toast.py', '--tournament-dir', str(ROOT / 'tournaments/Dunford'),
                    '--output-dir', folder]
            with patch('sys.argv', argv), patch('builtins.input', return_value=''), patch('sys.stdout', new=io.StringIO()):
                module['main'](self.cache)
            scripts = list(Path(folder).glob('*closing-ceremony.html'))
            self.assertEqual(len(scripts), 1)
            html = scripts[0].read_text(encoding='utf-8')
            for path in self.paths:
                roster = self.cache.table(path, 'Team and Program Information', 'OfficialTeamList')
                for name in roster['Team Name']:
                    self.assertIn(name, html)
            for handler in list(logging.getLogger('ceremony_generator').handlers):
                handler.close()
                logging.getLogger('ceremony_generator').removeHandler(handler)

    def test_single_division_warning_and_missing_summary(self):
        module = runpy.run_path(str(ROOT / 'fll-toast.py'))
        config = copy.deepcopy(self.config)
        config['INFO']['ojs_files'] = [dict(config['INFO']['ojs_files'][0], filename=self.paths[0])]
        config['INFO']['script_template'] = 'script.jinja'
        config['INFO']['summary_template'] = 'summary.jinja'
        for award in config['AWARDS']:
            if not award['DivAwd']:
                award['TournCount'] = 1
        # A real validator warning must not be treated as a template-variable string.
        df = self.cache.table(self.paths[0], 'Results and Rankings', 'TournamentData')
        df.loc[df['Advance?'] == 'Alt', 'Advance?'] = None
        self.cache.tables[(self.paths[0], 'Results and Rankings', 'TournamentData')] = df
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder)
            (base / 'tournament_config.json').write_text(json.dumps(config))
            (base / 'script.jinja').write_text('{{ team_list_D1 }} {{ div1_list }} {{ ADV_D1 }}')
            (base / 'summary.jinja').write_text('{{ team_list_D1 }}')
            argv = ['toast', '--tournament-dir', folder, '--output-dir', str(base / 'out')]
            try:
                with patch('sys.argv', argv), patch('builtins.input', return_value='y'), patch('sys.stdout', new=io.StringIO()):
                    module['main'](self.cache)
                outputs = list((base / 'out').glob('*.html'))
                self.assertEqual(len(outputs), 2)
                old = {p: p.read_bytes() for p in outputs}
                (base / 'summary.jinja').write_text('{{ nonexistent_winner }}')
                with patch('sys.argv', argv), patch('builtins.input', return_value='y'), patch('sys.stdout', new=io.StringIO()):
                    with self.assertRaises(ValueError):
                        module['main'](self.cache)
                self.assertEqual(old, {p: p.read_bytes() for p in outputs})
            finally:
                for handler in list(logging.getLogger('ceremony_generator').handlers):
                    handler.close()
                    logging.getLogger('ceremony_generator').removeHandler(handler)
